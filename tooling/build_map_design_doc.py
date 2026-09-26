from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "assets"
WORK = ROOT / "build" / "map_design_document"
OUTPUT = ROOT / "docs" / "Diseno_del_mapa_Energy_hasta_VOLT.docx"
WORK.mkdir(parents=True, exist_ok=True)
OUTPUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "17324D"
CYAN = "35BFD6"
PALE = "EAF5F8"
LIGHT = "F4F7F9"
GRID = "D9E2E7"
ORANGE = "F08A3C"
BLACK = RGBColor(0, 0, 0)


def font(size: int, bold: bool = False):
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


FONT_18 = font(18)
FONT_22 = font(22, True)
FONT_30 = font(30, True)


ROOMS = [
    {
        "id": "E01", "name": "Cámara de reactivación", "coord": "0, 0", "size": (1280, 720),
        "purpose": "Inicio seguro. Enseña carrera, salto y lectura del HUD sin presión.",
        "connections": "Izquierda cerrada; derecha a E02.",
        "progress": "Activa CP-01. No requiere habilidades ni permisos.",
        "geometry": [("PL-01", 0, 640, 1280, 80), ("PL-03", 320, 544, 128, 32), ("PL-04", 560, 480, 160, 24)],
        "objects": [("CP-01", 96, 500), ("M-0", 180, 570), ("DR-01", 1190, 528)],
        "notes": "Dejar una línea visual limpia hacia la salida. Sin enemigos. Usar señal de sector y luz cian.",
    },
    {
        "id": "E02", "name": "Distribuidor central", "coord": "1, 0", "size": (1280, 720),
        "purpose": "Hub de orientación. Presenta la rama superior, la inferior y la compuerta del jefe.",
        "connections": "Izquierda E01; arriba E03; abajo E06; derecha E09 mediante DR-02.",
        "progress": "DR-02 se abre desde E09 y permanece abierta. El acceso post-VOLT exige Dash.",
        "geometry": [("PL-01", 0, 640, 1280, 80), ("PL-03", 250, 544, 128, 32), ("PL-04", 520, 448, 160, 24), ("PL-06", 850, 340, 128, 32)],
        "objects": [("DR-01", 0, 528), ("DR-02", 1200, 528), ("SG-01", 520, 360), ("M-0", 150, 570)],
        "notes": "Desde una pasarela debe verse la silueta del reactor y la puerta hacia VOLT. Esta sala será el centro del minimapa.",
    },
    {
        "id": "E03", "name": "Ascensor de mantenimiento", "coord": "1, -1", "size": (1280, 720),
        "purpose": "Primera sala vertical. Enseña plataformas móviles y saltos encadenados.",
        "connections": "Abajo E02; derecha E04.",
        "progress": "Ruta opcional al principio. No concede progreso persistente.",
        "geometry": [("PL-01", 0, 640, 1280, 80), ("PL-02", 150, 540, 64, 32), ("PL-05", 430, 450, 128, 32), ("PL-04", 720, 350, 160, 24), ("PL-06", 990, 250, 128, 32)],
        "objects": [("EN-01", 750, 290), ("M-0", 120, 570), ("DR-01", 1190, 138)],
        "notes": "El ascenso debe poder completarse con salto normal. La caída devuelve a la zona inferior sin matar al jugador.",
    },
    {
        "id": "E04", "name": "Galería de turbinas", "coord": "2, -1", "size": (1280, 720),
        "purpose": "Combina movimiento horizontal, maquinaria y combate básico.",
        "connections": "Izquierda E03; derecha E05; bajada unilateral hacia E09.",
        "progress": "La bajada a E09 funciona como retorno corto después de explorar E05.",
        "geometry": [("PL-01", 0, 640, 340, 80), ("PL-01", 450, 640, 360, 80), ("PL-01", 920, 640, 360, 80), ("PL-05", 340, 520, 128, 32), ("PL-04", 620, 430, 160, 24), ("PL-03", 930, 340, 128, 32)],
        "objects": [("EN-01", 520, 570), ("EN-01", 1000, 570), ("SC-01", 700, 350), ("M-0", 100, 570)],
        "notes": "Colocar ventiladores y tuberías grandes al fondo. El conducto descendente debe ser visible pero no confundirse con una caída mortal.",
    },
    {
        "id": "E05", "name": "Centro de control", "coord": "3, -1", "size": (1280, 720),
        "purpose": "Entrega el permiso ENERGY-02 y conecta las dos alturas mediante ascensor.",
        "connections": "Izquierda E04; ascensor hacia E08; caída segura hacia E10 tras obtener el permiso.",
        "progress": "IT-02 solo concede ENERGY-02 cuando el interruptor de E08 está activo.",
        "geometry": [("PL-01", 0, 640, 1280, 80), ("PL-03", 260, 520, 128, 32), ("PL-04", 500, 420, 160, 24), ("PL-03", 760, 520, 128, 32), ("PL-05", 1060, 520, 128, 32)],
        "objects": [("IT-02", 580, 330), ("EN-02", 820, 440), ("EL-01", 1080, 450), ("M-0", 100, 570)],
        "notes": "Si falta energía, la consola muestra AUXILIAR OFF y señala E08. Con energía activa, concede el permiso y abre la compuerta inferior.",
    },
    {
        "id": "E06", "name": "Pozo de refrigeración", "coord": "1, 1", "size": (1280, 720),
        "purpose": "Descenso vertical controlado con plataformas unidireccionales.",
        "connections": "Arriba E02; derecha E07.",
        "progress": "Introduce superficies húmedas y electricidad sin bloquear el recorrido.",
        "geometry": [("PL-01", 0, 640, 1280, 80), ("PL-04", 110, 260, 160, 24), ("PL-04", 420, 380, 160, 24), ("PL-04", 730, 500, 160, 24), ("PL-02", 1040, 560, 64, 32)],
        "objects": [("HZ-02", 500, 600), ("M-0", 150, 190), ("DR-01", 1190, 528)],
        "notes": "El jugador debe ver la siguiente plataforma antes de saltar. Usar vapor y goteo como guía de profundidad.",
    },
    {
        "id": "E07", "name": "Conductos de refrigeración", "coord": "2, 1", "size": (1280, 720),
        "purpose": "Ruta de peligros ambientales y regreso vertical hacia E09.",
        "connections": "Izquierda E06; derecha E08; escalera superior a E09.",
        "progress": "Contiene un suelo agrietado que se abrirá más adelante con Golpe descendente.",
        "geometry": [("PL-01", 0, 640, 300, 80), ("PL-01", 410, 640, 320, 80), ("PL-01", 850, 640, 430, 80), ("PL-04", 280, 500, 160, 24), ("PL-07", 670, 430, 128, 32), ("PL-03", 970, 350, 128, 32)],
        "objects": [("HZ-01", 300, 590), ("HZ-02", 730, 600), ("EN-01", 500, 570), ("BR-01", 900, 610), ("M-0", 90, 570)],
        "notes": "La barrera BR-01 debe verse distinta del suelo normal. No es necesaria para llegar a VOLT; queda como promesa de regreso.",
    },
    {
        "id": "E08", "name": "Banco de capacitores", "coord": "3, 1", "size": (1280, 720),
        "purpose": "Objetivo de la rama inferior. Restablece la energía auxiliar.",
        "connections": "Izquierda E07; ascensor a E05; conexión técnica a E10.",
        "progress": "IT-01 activa AUX_POWER y desbloquea el ascensor E08-E05.",
        "geometry": [("PL-01", 0, 640, 1280, 80), ("PL-03", 220, 540, 128, 32), ("PL-04", 470, 450, 160, 24), ("PL-04", 760, 360, 160, 24), ("PL-05", 1040, 520, 128, 32)],
        "objects": [("EN-02", 520, 380), ("IT-01", 860, 570), ("EL-01", 1060, 450), ("M-0", 100, 570)],
        "notes": "Tras activar el interruptor, cambiar luces naranjas a cian y encender generadores. La transformación visual confirma el progreso.",
    },
    {
        "id": "E09", "name": "Relé central", "coord": "2, 0", "size": (1280, 720),
        "purpose": "Sala de convergencia y primer atajo importante.",
        "connections": "Arriba E04; abajo E07; izquierda E02; derecha E10.",
        "progress": "Abre DR-02 hacia E02. La salida a E10 requiere AUX_POWER y ENERGY-02.",
        "geometry": [("PL-01", 0, 640, 1280, 80), ("PL-03", 240, 520, 128, 32), ("PL-04", 500, 420, 160, 24), ("PL-03", 780, 520, 128, 32), ("PL-06", 980, 330, 128, 32)],
        "objects": [("DR-02", 0, 528), ("EN-02", 610, 350), ("EN-01", 820, 570), ("DR-02", 1200, 528), ("M-0", 120, 570)],
        "notes": "La batalla debe usar máximo dos enemigos activos. Al terminar, abrir el acceso a E02 y mostrar una señal hacia VOLT.",
    },
    {
        "id": "E10", "name": "Corredor de alto voltaje", "coord": "3, 0", "size": (1280, 720),
        "purpose": "Examen previo al jefe con electricidad, plataforma móvil y enemigo aéreo.",
        "connections": "Izquierda E09; derecha E11; túnel de Dash post-VOLT hacia E02.",
        "progress": "El túnel de Dash es visible, pero permanece inaccesible antes del jefe.",
        "geometry": [("PL-01", 0, 640, 270, 80), ("PL-01", 420, 640, 300, 80), ("PL-01", 890, 640, 390, 80), ("PL-05", 280, 520, 128, 32), ("PL-04", 610, 420, 160, 24), ("PL-06", 930, 330, 128, 32)],
        "objects": [("HZ-02", 270, 600), ("HZ-02", 720, 600), ("EN-03", 700, 280), ("DG-01", 1040, 540), ("M-0", 80, 570)],
        "notes": "Reducir la densidad si los controles quedan tapados. El Drone debe atacar cuando el jugador ya está sobre una superficie estable.",
    },
    {
        "id": "E11", "name": "Cámara de aislamiento", "coord": "4, 0", "size": (1280, 720),
        "purpose": "Descanso, checkpoint y anticipación narrativa antes de VOLT.",
        "connections": "Izquierda E10; derecha E12.",
        "progress": "Activa CP-02. La puerta de la arena se cierra al iniciar el combate.",
        "geometry": [("PL-01", 0, 640, 1280, 80), ("PL-03", 300, 540, 128, 32), ("PL-03", 850, 540, 128, 32)],
        "objects": [("CP-01", 170, 500), ("SG-01", 540, 400), ("DR-01", 1190, 528), ("M-0", 350, 570)],
        "notes": "Sin enemigos. La música y el ruido industrial deben bajar. Mostrar el texto NÚCLEO DE SEGURIDAD VOLT tras activar el checkpoint.",
    },
    {
        "id": "E12", "name": "Núcleo de VOLT", "coord": "5, 0", "size": (1920, 720),
        "purpose": "Arena del primer jefe, de una pantalla y media de ancho.",
        "connections": "Izquierda E11; salida postcombate hacia el corredor de Dash.",
        "progress": "Derrotar a VOLT guarda bossDefeated=volt y concede el módulo Dash.",
        "geometry": [("PL-01", 0, 640, 1920, 80), ("PL-03", 260, 520, 160, 32), ("PL-03", 760, 520, 160, 32), ("PL-03", 1260, 520, 160, 32), ("PL-03", 1580, 520, 160, 32)],
        "objects": [("BOSS", 980, 430), ("HZ-03", 400, 600), ("HZ-03", 900, 600), ("HZ-03", 1400, 600), ("M-0", 180, 570)],
        "notes": "Dividir el suelo en cuatro sectores para Grid Surge. Mantener dos plataformas laterales seguras y espacio para saltar Ground Slam.",
    },
]


ASSET_REFS = [
    ("PL-01", "Suelo modular", "environment/tile_normal.png", "Piso y paredes sólidas; repetir en una cuadrícula de 32 px."),
    ("PL-02", "Plataforma pequeña", "environment/platform_small.png", "Apoyo corto para saltos de precisión."),
    ("PL-03", "Plataforma mediana", "environment/platform_medium.png", "Apoyo estándar y descanso entre saltos."),
    ("PL-04", "Plataforma completa", "environment/platform_full.png", "Plataforma unidireccional atravesable desde abajo."),
    ("PL-05", "Plataforma móvil", "environment/moving_platform_center.png", "Movimiento horizontal o vertical con pausa en extremos."),
    ("PL-06", "Plataforma suspendida", "environment/platform_suspended.png", "Apoyo elevado y lectura vertical del escenario."),
    ("PL-07", "Plataforma de rejilla", "environment/platform_grate.png", "Variante visual para conductos y refrigeración."),
    ("PL-08", "Plataforma con soporte", "environment/platform_support.png", "Marca plataformas principales y cambios de altura."),
    ("HZ-01", "Pinchos", "environment/spikes_active.png", "Daño por contacto con espacio seguro antes y después."),
    ("HZ-02", "Suelo electrificado", "environment/electric_floor_active.png", "Peligro temporal o permanente con señal previa."),
    ("HZ-03", "Descarga por sectores", "environment/surface_discharge.png", "Ataque de VOLT; activar por segmentos de la arena."),
    ("DG-01", "Barrera de Dash", "environment/laser_active.png", "Ruta visible antes de VOLT y transitable al obtener Dash."),
    ("DR-01", "Puerta normal", "environment/door_closed.png", "Cambio de sala sin requisito especial."),
    ("DR-02", "Puerta bloqueada", "environment/door_blocked.png", "Comunica permiso, interruptor o atajo pendiente."),
    ("IT-01", "Interruptor de energía", "environment/switch_active.png", "Activa AUX_POWER y el ascensor E08-E05."),
    ("IT-02", "Consola de permiso", "environment/control_console.png", "Entrega ENERGY-02 cuando la energía auxiliar está activa."),
    ("CP-01", "Checkpoint", "environment/checkpoint_active.png", "Cura, guarda y fija el punto de reaparición."),
    ("EL-01", "Cabina de ascensor", "environment/moving_platform_center.png", "Usar con pilares y guías verticales; conexión E05-E08."),
    ("BR-01", "Suelo agrietado", "environment/wear_crack.png", "Ruta futura que requerirá Golpe descendente."),
    ("SC-01", "Señal Energy", "environment/sign_energy.png", "Orienta hacia objetivos y refuerza el lenguaje del sector."),
    ("EN-01", "Patrol Unit", "sprites/patrol/idle/frame_00.png", "Enemigo terrestre para pasarelas amplias."),
    ("EN-02", "Watcher", "sprites/watcher/idle/frame_00.png", "Torre para controlar plataformas y cruces verticales."),
    ("EN-03", "Drone", "sprites/drone/idle/frame_00.png", "Amenaza aérea en salas con espacio para maniobrar."),
    ("M-0", "Personaje", "sprites/m0/idle/frame_00.png", "Referencia de escala: collider aproximado 36 x 58 px."),
    ("BOSS", "VOLT", "sprites/volt/idle/frame_00.png", "Jefe del sector; reservar un área limpia alrededor."),
    ("DC-01", "Generador principal", "environment/generator_main.png", "Decoración de gran tamaño para E08 y E12."),
    ("DC-02", "Cables activos", "environment/active_cables.png", "Guía visual de energía activa y conexiones entre máquinas."),
    ("DC-03", "Luz industrial", "environment/industrial_lamp.png", "Puntos de interés, entradas y superficies seguras."),
]


ICON_PATHS = {code: ASSETS / rel for code, _, rel, _ in ASSET_REFS}


def paste_icon(canvas: Image.Image, code: str, x: int, y: int, max_size=(72, 72)):
    path = ICON_PATHS.get(code)
    if path is None or not path.exists():
        return
    icon = Image.open(path).convert("RGBA")
    icon.thumbnail(max_size, Image.Resampling.NEAREST)
    canvas.alpha_composite(icon, (int(x - icon.width / 2), int(y - icon.height)))


def make_map_diagram() -> Path:
    out = WORK / "mapa_general.png"
    image = Image.new("RGB", (1600, 900), "#F4F7F9")
    draw = ImageDraw.Draw(image)
    positions = {
        "E01": (70, 390), "E02": (290, 390), "E09": (730, 390), "E10": (950, 390),
        "E11": (1170, 390), "E12": (1390, 390), "E03": (290, 120), "E04": (510, 120),
        "E05": (730, 120), "E06": (290, 660), "E07": (510, 660), "E08": (730, 660),
    }
    links = [
        ("E01", "E02", "normal"), ("E02", "E03", "normal"), ("E03", "E04", "normal"),
        ("E04", "E05", "normal"), ("E02", "E06", "normal"), ("E06", "E07", "normal"),
        ("E07", "E08", "normal"), ("E04", "E09", "normal"), ("E07", "E09", "normal"),
        ("E05", "E08", "elevator"), ("E02", "E09", "locked"), ("E09", "E10", "gate"),
        ("E10", "E11", "normal"), ("E11", "E12", "boss"),
    ]
    for a, b, kind in links:
        ax, ay = positions[a]; bx, by = positions[b]
        color = "#2D8FA3" if kind in {"normal", "elevator"} else "#E77D32"
        width = 10 if kind in {"gate", "boss"} else 7
        draw.line((ax + 80, ay + 46, bx + 80, by + 46), fill=color, width=width)
        if kind == "locked":
            mx, my = (ax + bx) / 2 + 80, (ay + by) / 2 + 46
            draw.ellipse((mx - 15, my - 15, mx + 15, my + 15), fill="#F4F7F9", outline=color, width=5)
    for room in ROOMS:
        x, y = positions[room["id"]]
        w = 180 if room["id"] != "E12" else 195
        fill = "#17324D" if room["id"] != "E12" else "#7A2D2D"
        draw.rounded_rectangle((x, y, x + w, y + 92), radius=16, fill=fill, outline="#35BFD6", width=4)
        draw.text((x + 16, y + 12), room["id"], font=FONT_30, fill="white")
        label = room["name"].replace("Cámara de ", "").replace("Conductos de ", "")
        draw.text((x + 16, y + 56), label[:20], font=FONT_18, fill="#CFEAF0")
    draw.text((70, 28), "SECTOR ENERGY  MAPA HASTA VOLT", font=FONT_30, fill="#17324D")
    draw.text((70, 820), "Azul: conexión normal  |  Naranja: bloqueo o requisito  |  E05-E08: ascensor vertical", font=FONT_22, fill="#304C5C")
    image.save(out)
    return out


PLATFORM_COLORS = {
    "PL-01": "#3A5261", "PL-02": "#65869A", "PL-03": "#65869A", "PL-04": "#8DB6C3",
    "PL-05": "#37A7C4", "PL-06": "#7997A7", "PL-07": "#63808D", "PL-08": "#536E7D",
}


def make_room_diagram(room: dict) -> Path:
    out = WORK / f'{room["id"]}.png'
    ww, wh = room["size"]
    image = Image.new("RGBA", (960, 540), "#091721")
    draw = ImageDraw.Draw(image)
    sx, sy = 960 / ww, 540 / wh
    for gx in range(0, ww + 1, 160):
        x = int(gx * sx); draw.line((x, 0, x, 540), fill="#153040", width=1)
    for gy in range(0, wh + 1, 120):
        y = int(gy * sy); draw.line((0, y, 960, y), fill="#153040", width=1)
    for code, x, y, w, h in room["geometry"]:
        rect = (int(x * sx), int(y * sy), int((x + w) * sx), int((y + h) * sy))
        draw.rectangle(rect, fill=PLATFORM_COLORS.get(code, "#586D79"), outline="#B9D4DE", width=2)
        if w * sx > 55:
            draw.text((rect[0] + 5, max(2, rect[1] + 3)), code, font=FONT_18, fill="white")
    for code, x, y in room["objects"]:
        px, py = int(x * sx), int(y * sy)
        paste_icon(image, code, px, py, max_size=(70, 70))
        draw.text((px + 10, max(4, py - 70)), code, font=FONT_18, fill="#FFD07A")
    draw.rectangle((0, 0, 959, 539), outline="#35BFD6", width=4)
    draw.text((20, 18), f'{room["id"]}  {room["name"]}', font=FONT_30, fill="#E7F7FB")
    draw.text((20, 58), f'Mundo {ww} x {wh} px  |  Coordenada de mapa {room["coord"]}', font=FONT_18, fill="#A9C8D2")
    image.convert("RGB").save(out, quality=95)
    return out


def make_asset_card(code: str, name: str, rel: str) -> Path:
    out = WORK / f"ref_{code.replace('-', '_')}.png"
    bg = Image.new("RGBA", (520, 280), "#0B1A25")
    draw = ImageDraw.Draw(bg)
    path = ASSETS / rel
    art = Image.open(path).convert("RGBA")
    art.thumbnail((430, 185), Image.Resampling.NEAREST)
    bg.alpha_composite(art, ((520 - art.width) // 2, 68 + (175 - art.height) // 2))
    draw.rectangle((1, 1, 518, 278), outline="#315568", width=3)
    draw.text((20, 16), f"{code}  {name}", font=FONT_22, fill="#E8F6FA")
    bg.convert("RGB").save(out, quality=95)
    return out


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRID, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_run(run, size=10.5, bold=False, color=BLACK):
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def add_label_paragraph(cell, label: str, value: str, size=9.2):
    p = cell.add_paragraph() if cell.text else cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.02
    r = p.add_run(label + " ")
    style_run(r, size=size, bold=True, color=RGBColor(23, 50, 77))
    r = p.add_run(value)
    style_run(r, size=size)


def add_table(doc, headers, rows, widths=None, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, NAVY); set_cell_border(cell); set_cell_margins(cell)
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header); style_run(r, size=font_size, bold=True, color=RGBColor(255, 255, 255))
        if widths: cell.width = Inches(widths[i])
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            set_cell_shading(cell, "FFFFFF" if ri % 2 == 0 else PALE)
            set_cell_border(cell); set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths: cell.width = Inches(widths[i])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value)); style_run(r, size=font_size, bold=(i == 0))
    return table


def add_heading(doc, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(5)
    for r in p.runs:
        style_run(r, size=17 if level == 1 else 13, bold=True, color=BLACK)
    return p


def add_body(doc, text: str, bold_lead: str | None = None, size=10.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    if bold_lead:
        r = p.add_run(bold_lead); style_run(r, size=size, bold=True)
    r = p.add_run(text); style_run(r, size=size)
    return p


def add_room_card(doc: Document, room: dict, diagram: Path):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_border(cell, "AFC4CE", "8"); set_cell_margins(cell, 100, 120, 100, 120)
    set_cell_shading(cell, LIGHT)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f'{room["id"]}  {room["name"]}')
    style_run(r, size=13.5, bold=True, color=RGBColor(23, 50, 77))
    p = cell.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(diagram), width=Inches(6.55))
    add_label_paragraph(cell, "Objetivo:", room["purpose"])
    add_label_paragraph(cell, "Conexiones:", room["connections"])
    add_label_paragraph(cell, "Progreso:", room["progress"])
    geo = "; ".join(f"{c} ({x},{y},{w},{h})" for c, x, y, w, h in room["geometry"])
    add_label_paragraph(cell, "Geometría x y ancho alto:", geo, size=8.5)
    objs = "; ".join(f"{c} ({x},{y})" for c, x, y in room["objects"])
    add_label_paragraph(cell, "Objetos y entidades:", objs, size=8.5)
    add_label_paragraph(cell, "Dirección artística:", room["notes"])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def build_document():
    map_path = make_map_diagram()
    room_diagrams = {room["id"]: make_room_diagram(room) for room in ROOMS}
    asset_cards = {code: make_asset_card(code, name, rel) for code, name, rel, _ in ASSET_REFS}

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = Inches(0.65); section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7); section.right_margin = Inches(0.7)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"; styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Aptos Display"; styles["Title"].font.size = Pt(28); styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = BLACK
    title_ppr = styles["Title"].element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    for style_name in ("Heading 1", "Heading 2"):
        styles[style_name].font.name = "Aptos Display"; styles[style_name].font.color.rgb = BLACK

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    title.add_run("Diseño del mapa Energy hasta VOLT")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    r = subtitle.add_run("Guía de salas plataformas conexiones y referencias visuales")
    style_run(r, size=13, color=RGBColor(70, 91, 103))
    add_body(doc, "Este documento define una propuesta completa para el primer sector de Sistema Caído. El recorrido usa doce salas distribuidas en tres alturas, dos ramas que se conectan entre sí, un atajo central y una ruta de regreso que se habilita después de derrotar a VOLT.")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(map_path), width=Inches(7.05))
    cap = doc.add_paragraph("Plano general del sector Energy")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in cap.runs: style_run(r, size=9, bold=True, color=RGBColor(70, 91, 103))

    doc.add_page_break()
    add_heading(doc, "Recorrido y progresión", 1)
    add_body(doc, "La ruta recomendada comienza en E01, presenta el hub E02 y conduce primero a la rama inferior para restablecer la energía. El ascensor conecta E08 con E05, donde se obtiene ENERGY-02. Las dos ramas convergen en E09; desde allí se abre el atajo a E02 y continúa la aproximación a VOLT.")
    add_table(doc, ["Paso", "Ruta", "Resultado"], [
        ("1", "E01 → E02", "Tutorial, checkpoint y presentación del hub."),
        ("2", "E02 → E06 → E07 → E08", "Restablecer AUX_POWER."),
        ("3", "E08 → E05", "Activar ascensor y obtener ENERGY-02."),
        ("4", "E05/E04 → E09", "Abrir el atajo permanente hacia E02."),
        ("5", "E09 → E10 → E11", "Prueba final y checkpoint del jefe."),
        ("6", "E11 → E12", "Derrotar a VOLT y obtener Dash."),
    ], widths=[0.6, 2.3, 3.7], font_size=9.5)
    add_heading(doc, "Reglas de construcción", 1)
    rules = [
        "Usar salas de 1280 x 720 px y una cuadrícula base de 32 px. La arena E12 puede medir 1920 x 720 px.",
        "Mantener 96 px libres alrededor de puertas y checkpoints. Evitar peligros bajo los controles táctiles en los extremos inferiores.",
        "Toda plataforma obligatoria debe poder alcanzarse con las habilidades disponibles antes de VOLT: carrera, salto variable y ataque.",
        "Las salidas verticales deben mostrar visualmente el siguiente destino. No usar caídas ciegas como conexión principal.",
        "Los bloqueos futuros de Dash, Golpe descendente y Doble salto se muestran, pero no forman parte de la ruta obligatoria.",
        "Limitar a tres enemigos simultáneos por cámara. Combinar máximo dos tipos de enemigo antes del jefe.",
    ]
    for item in rules:
        p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item); style_run(r, size=10.2)
    add_heading(doc, "Convención de coordenadas", 2)
    add_body(doc, "Las coordenadas de cada ficha usan el origen 0,0 en la esquina superior izquierda. Los valores se expresan como x, y, ancho y alto. Son medidas sugeridas para bocetar y pueden ajustarse durante pruebas de movimiento.")

    for room in ROOMS:
        doc.add_page_break()
        add_heading(doc, f"Sala {room['id']}", 1)
        add_room_card(doc, room, room_diagrams[room["id"]])

    groups = [
        ("Anexo de plataformas", ASSET_REFS[0:8]),
        ("Anexo de peligros e interactivos", ASSET_REFS[8:16]),
        ("Anexo de peligros e interactivos continuación", ASSET_REFS[16:20]),
        ("Anexo de personajes y decoración", ASSET_REFS[20:]),
    ]
    for group_index, (heading, items) in enumerate(groups):
        if group_index in {0, 3}:
            doc.add_page_break()
        add_heading(doc, heading, 1)
        add_body(doc, "Los códigos de este anexo coinciden con las etiquetas utilizadas en los planos de sala. Las imágenes son referencias ya integradas en el proyecto y pueden usarse como base para construir el escenario final.")
        table = doc.add_table(rows=0, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for i in range(0, len(items), 2):
            cells = table.add_row().cells
            for ci in range(2):
                cell = cells[ci]; cell.width = Inches(3.35); cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
                set_cell_border(cell, "C7D5DB", "7"); set_cell_margins(cell, 90, 100, 100, 100)
                set_cell_shading(cell, "FFFFFF" if (i // 2) % 2 == 0 else LIGHT)
                if i + ci >= len(items):
                    continue
                code, name, _, usage = items[i + ci]
                p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(asset_cards[code]), width=Inches(2.78))
                p = cell.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.space_after = Pt(2)
                r = p.add_run(f"{code}  {name}"); style_run(r, size=10, bold=True, color=RGBColor(23, 50, 77))
                p = cell.add_paragraph(); p.paragraph_format.space_after = Pt(0)
                r = p.add_run(usage); style_run(r, size=8.7)

    add_heading(doc, "Resumen de producción", 1)
    add_body(doc, "Para diseñar el mapa, conviene trabajar primero con bloques y colisiones, después colocar puertas e interactivos, luego enemigos y finalmente decoración. La composición visual nunca debe alterar una distancia de salto ya validada.")
    add_table(doc, ["Prioridad", "Tarea", "Resultado esperado"], [
        ("1", "Construir E01, E02 y conexiones verticales", "Validar orientación y escala general."),
        ("2", "Construir rama inferior E06-E08", "Validar descenso, peligros y AUX_POWER."),
        ("3", "Construir rama superior E03-E05", "Validar ascenso, ascensor y ENERGY-02."),
        ("4", "Construir convergencia E09-E11", "Validar atajo, retorno y dificultad previa."),
        ("5", "Construir E12", "Validar espacio y patrones de VOLT."),
        ("6", "Añadir decoración y señales", "Mejorar lectura sin cambiar colisiones."),
    ], widths=[0.8, 2.8, 3.2], font_size=9.4)
    add_heading(doc, "Lista mínima de recursos", 2)
    add_body(doc, "8 familias de plataformas, 3 peligros, 2 tipos de puerta, interruptor, consola, checkpoint, ascensor, 3 enemigos, M-0, VOLT, señales, generadores, cables y luces industriales. Todos aparecen identificados en los anexos.")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
